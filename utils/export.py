# utils/export.py
# PDF and Excel export utilities

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import os

# ─── PDF EXPORT ───────────────────────────────────────────────────────────────

def export_transactions_pdf(transactions, period, output_path):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A4),
        rightMargin=0.5*inch, leftMargin=0.5*inch,
        topMargin=0.5*inch, bottomMargin=0.5*inch
    )

    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = ParagraphStyle('title', fontSize=16, fontName='Helvetica-Bold',
                                  alignment=TA_CENTER, spaceAfter=6)
    sub_style = ParagraphStyle('sub', fontSize=10, fontName='Helvetica',
                                alignment=TA_CENTER, spaceAfter=16,
                                textColor=colors.grey)

    elements.append(Paragraph("ParkEase — Transaction Report", title_style))
    elements.append(Paragraph(
        f"Period: {period.title()}  |  Generated: {datetime.now().strftime('%B %d, %Y %H:%M')}",
        sub_style
    ))

    # Table headers
    headers = ['Txn ID', 'Plate', 'Type', 'Date', 'Time In', 'Time Out',
               'Duration', 'Charge (₱)', 'Status', 'Staff (Entry)', 'Staff (Exit)']

    rows = [headers]
    total = 0.0
    for t in transactions:
        ti = t['time_in']
        to = t['time_out']
        ti_s = ti.strftime("%H:%M:%S") if isinstance(ti, datetime) else str(ti)
        to_s = to.strftime("%H:%M:%S") if isinstance(to, datetime) and to else '—'
        di = t['date_in']
        di_s = di.strftime("%Y-%m-%d") if hasattr(di, 'strftime') else str(di)
        mins = t.get('duration_minutes') or 0
        h, m = divmod(mins, 60)
        charge = float(t.get('total_charge') or 0)
        total += charge
        rows.append([
            t['transaction_id'], t['plate_number'], t['vehicle_type'],
            di_s, ti_s, to_s,
            f"{h}h {m}m" if mins else '—',
            f"{charge:.2f}",
            t.get('status', '—').upper(),
            t.get('staff_entry') or '—',
            t.get('staff_exit') or '—'
        ])

    # Total row
    rows.append(['', '', '', '', '', '', 'TOTAL', f"{total:.2f}", '', '', ''])

    col_widths = [1.4*inch, 0.9*inch, 0.9*inch, 0.9*inch, 0.8*inch, 0.8*inch,
                  0.7*inch, 0.8*inch, 0.8*inch, 1.1*inch, 1.1*inch]

    table = Table(rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        # Header
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#3B82F6')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,0), 9),
        ('ALIGN', (0,0), (-1,0), 'CENTER'),
        ('BOTTOMPADDING', (0,0), (-1,0), 8),
        ('TOPPADDING', (0,0), (-1,0), 8),
        # Body
        ('FONTSIZE', (0,1), (-1,-2), 8),
        ('ROWBACKGROUNDS', (0,1), (-1,-2), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('ALIGN', (7,1), (7,-1), 'RIGHT'),
        ('TOPPADDING', (0,1), (-1,-1), 6),
        ('BOTTOMPADDING', (0,1), (-1,-1), 6),
        # Total row
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#F1F5F9')),
        ('FONTNAME', (6,-1), (7,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,-1), (-1,-1), 9),
    ]))

    elements.append(table)
    doc.build(elements)
    return output_path


def export_receipt_pdf(txn_data, output_path):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=2*inch, leftMargin=2*inch,
        topMargin=1*inch, bottomMargin=1*inch
    )

    styles = getSampleStyleSheet()
    elements = []

    center = ParagraphStyle('c', alignment=TA_CENTER, spaceAfter=4)
    bold_center = ParagraphStyle('bc', alignment=TA_CENTER,
                                  fontName='Helvetica-Bold', fontSize=14, spaceAfter=8)

    elements.append(Paragraph("ParkEase", bold_center))
    elements.append(Paragraph("Parking Receipt", ParagraphStyle(
        'sub', alignment=TA_CENTER, fontSize=11, spaceAfter=4)))
    elements.append(Paragraph(
        datetime.now().strftime("%B %d, %Y  %H:%M"),
        ParagraphStyle('date', alignment=TA_CENTER, fontSize=9,
                       textColor=colors.grey, spaceAfter=20)))

    # Receipt details
    data = [
        ['Transaction ID', txn_data.get('transaction_id', '—')],
        ['Plate Number',   txn_data.get('plate_number', '—')],
        ['Vehicle Type',   txn_data.get('vehicle_type', '—')],
        ['Time In',        str(txn_data.get('time_in', '—'))],
        ['Time Out',       str(txn_data.get('time_out', '—'))],
        ['Duration',       f"{txn_data.get('duration_minutes', 0) // 60}h "
                           f"{txn_data.get('duration_minutes', 0) % 60}m"],
        ['', ''],
        ['TOTAL CHARGE',   f"₱{txn_data.get('total_charge', 0):.2f}"],
    ]

    t = Table(data, colWidths=[2.5*inch, 2.5*inch])
    t.setStyle(TableStyle([
        ('FONTSIZE', (0,0), (-1,-1), 10),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('TEXTCOLOR', (0,0), (0,-1), colors.grey),
        ('FONTNAME', (0,-1), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,-1), (-1,-1), 13),
        ('TEXTCOLOR', (1,-1), (1,-1), colors.HexColor('#10B981')),
        ('LINEABOVE', (0,-2), (-1,-2), 1, colors.HexColor('#E2E8F0')),
        ('LINEABOVE', (0,-1), (-1,-1), 1, colors.HexColor('#3B82F6')),
    ]))

    elements.append(t)
    elements.append(Spacer(1, 20))
    elements.append(Paragraph("Thank you for parking with us!",
                               ParagraphStyle('thanks', alignment=TA_CENTER,
                                              fontSize=10, textColor=colors.grey)))
    doc.build(elements)
    return output_path


# ─── EXCEL EXPORT ─────────────────────────────────────────────────────────────

def export_transactions_excel(transactions, period, output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"

    # Styles
    header_fill = PatternFill("solid", fgColor="3B82F6")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    title_font  = Font(bold=True, size=14)
    total_font  = Font(bold=True, size=11)
    alt_fill    = PatternFill("solid", fgColor="F8FAFC")
    center      = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title
    ws.merge_cells('A1:K1')
    ws['A1'] = "ParkEase — Transaction Report"
    ws['A1'].font = title_font
    ws['A1'].alignment = center

    ws.merge_cells('A2:K2')
    ws['A2'] = f"Period: {period.title()}  |  Generated: {datetime.now().strftime('%B %d, %Y %H:%M')}"
    ws['A2'].alignment = center
    ws['A2'].font = Font(color="94A3B8", size=10)

    # Headers
    headers = ['Txn ID', 'Plate', 'Type', 'Date', 'Time In', 'Time Out',
               'Duration', 'Charge (₱)', 'Status', 'Staff (Entry)', 'Staff (Exit)']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = thin_border

    # Data
    total = 0.0
    for row_idx, t in enumerate(transactions, 5):
        ti = t['time_in']
        to = t['time_out']
        ti_s = ti.strftime("%H:%M:%S") if isinstance(ti, datetime) else str(ti)
        to_s = to.strftime("%H:%M:%S") if isinstance(to, datetime) and to else '—'
        di = t['date_in']
        di_s = di.strftime("%Y-%m-%d") if hasattr(di, 'strftime') else str(di)
        mins = t.get('duration_minutes') or 0
        h, m = divmod(mins, 60)
        charge = float(t.get('total_charge') or 0)
        total += charge

        row_data = [
            t['transaction_id'], t['plate_number'], t['vehicle_type'],
            di_s, ti_s, to_s,
            f"{h}h {m}m" if mins else '—',
            charge,
            t.get('status', '—').upper(),
            t.get('staff_entry') or '—',
            t.get('staff_exit') or '—'
        ]

        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
            if fill:
                cell.fill = fill

    # Total row
    total_row = len(transactions) + 5
    ws.cell(row=total_row, column=7, value='TOTAL').font = total_font
    total_cell = ws.cell(row=total_row, column=8, value=total)
    total_cell.font = Font(bold=True, size=11, color="10B981")
    total_cell.number_format = '₱#,##0.00'

    # Column widths
    col_widths = [22, 12, 12, 12, 12, 12, 10, 12, 12, 18, 18]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.row_dimensions[4].height = 30
    ws.freeze_panes = 'A5'

    wb.save(output_path)
    return output_path